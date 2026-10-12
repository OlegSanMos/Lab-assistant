from pathlib import Path
from urllib.parse import quote

from docx import Document
from PIL import Image

from conftest import PROJECT_DIR, ScriptedChatModel, course_responder
from labmentor.assistant import LabAssistant, Options
from labmentor.config import Settings
from labmentor.data_analysis import fit_polynomial, load_table
from labmentor.graphics import plot_fit
from labmentor.report import compose_report, export_report_docx, figure_catalog
from labmentor.ui import served

MANUAL_FIGURE = "images/lab_04_cmos/fig01.png"
DRAFT = """## Краткие теоретические сведения
Простейшей КМОП-схемой является инвертор (см. схему на рис. 3).

![Рис. 3. Электрическая схема КМОП-пары (инвертора)](images/lab/fig03.png)

*Рис. 3. Электрическая схема КМОП-пары (инвертора)*

Передаточная характеристика показана на рис. 5.

![Рис. 5. Передаточная характеристика](images/lab/missing.png)

## Результаты измерений и их обработка
Коэффициент найден по графику.

## Выводы
Схема на рис. 3 работает.
"""


def picture(path: Path, size=(400, 300), color="red") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)
    return path


def finder(root: Path):
    def find(ref):
        path = Path(ref) if Path(ref).is_absolute() else root / ref
        return path if path.is_file() else None

    return find


def measured(tmp_path, sample):
    df = load_table(sample("diode_iv.csv"))
    return df, plot_fit(df, fit_polynomial(df, "U_В", "I_мА", 2, 0.95), tmp_path)


def test_figures_are_numbered_in_order_and_references_follow(tmp_path, sample):
    manuals = tmp_path / "manuals"
    scheme = picture(manuals / "images/lab/fig03.png")
    photo = picture(tmp_path / "photo.webp", (300, 600), "blue")
    df, plot = measured(tmp_path, sample)
    final = compose_report(DRAFT, df, plot, [(str(photo), "Фото установки")], finder(manuals))
    assert "(см. схему на рис. 1)" in final and "Схема на рис. 1 работает" in final
    assert "*Рис. 3." not in final and "[вставить рисунок: Передаточная характеристика]" in final
    assert f"![Рисунок 1. Электрическая схема КМОП-пары (инвертора)](<{scheme}>)" in final
    assert "показаны на рисунке 2." in final
    assert "Материалы, приложенные к работе, приведены на рисунке 3." in final
    order = ["Таблица 1 – Результаты измерений", "Рисунок 2.", "Рисунок 3.", "Коэффициент найден"]
    assert [final.index(text) for text in order] == sorted(final.index(text) for text in order)


def test_manual_figures_mentioned_in_text_are_added(tmp_path):
    manuals = tmp_path / "manuals"
    for name in ("fig02", "fig05", "fig08"):
        picture(manuals / f"images/lab/{name}.png")
    catalog = figure_catalog(
        "Текст методички\n![Рис. 2. Передаточная характеристика](images/lab/fig02.png)\n"
        "![Рис. 5. Логические схемы](images/lab/fig05.png)\n"
        "![Рис. 8. Схема для измерения](images/lab/fig08.png)\n![Фото](images/lab/x.png)"
    )
    assert catalog == {
        2: ("Передаточная характеристика", "images/lab/fig02.png"),
        5: ("Логические схемы", "images/lab/fig05.png"),
        8: ("Схема для измерения", "images/lab/fig08.png"),
    }
    draft = (
        "## Ход работы\n1. Собрать схему по рис. 8.\n2. Снять характеристику (рис. 2, а).\n\n"
        "*Рис. 5. Логические схемы*\n\n## Выводы\n"
        "Схема на рис. 8 работает, рис. 9 в методичке нет.\n"
    )
    final = compose_report(draft, find_image=finder(manuals), manual=catalog)
    assert "1. Собрать схему по рис. 1." in final and "(рис. 2, а)" in final
    assert "Схема на рис. 1 работает, рис. 9 в методичке нет." in final
    names = ["Рисунок 1. Схема для измерения", "Рисунок 2. Передаточная", "Рисунок 3. Логические"]
    assert [final.index(name) for name in names] == sorted(final.index(name) for name in names)
    assert final.index("2. Снять") < final.index(names[0]) and "*Рис. 5" not in final
    assert compose_report(draft, find_image=finder(manuals)).count("![") == 0


def test_docx_contains_real_pictures_with_captions(tmp_path, sample):
    manuals = tmp_path / "manuals"
    picture(manuals / "images/lab/fig03.png")
    photo = picture(tmp_path / "photo.webp", (300, 600), "blue")
    df, plot = measured(tmp_path, sample)
    path = export_report_docx(
        DRAFT,
        "ЛР 1.4",
        tmp_path / "out",
        df,
        plot,
        [(str(photo), "Фото установки")],
        finder(manuals),
    )
    document = Document(path)
    texts = [p.text for p in document.paragraphs]
    assert [t for t in texts if t.startswith("Рисунок ")] == [
        "Рисунок 1. Электрическая схема КМОП-пары (инвертора)",
        "Рисунок 2. Экспериментальная зависимость и аппроксимация МНК",
        "Рисунок 3. Фото установки",
    ]
    assert len(document.inline_shapes) == 3
    assert not any("![" in text or "](" in text for text in texts)
    portrait = document.inline_shapes[2]
    assert round(portrait.width.cm, 1) == 6.0 and round(portrait.height.cm, 1) == 12.0


def test_assistant_finds_only_known_images(tmp_path):
    settings = Settings(
        provider="openai", openai_api_key="k", embeddings_provider="none", storage_dir=tmp_path
    )
    assistant = LabAssistant(settings, llm_factory=lambda **kwargs: None)
    figure = PROJECT_DIR / "data" / "manuals" / MANUAL_FIGURE
    assert assistant.find_image(MANUAL_FIGURE) == figure
    assert assistant.find_image(f"data/manuals/{MANUAL_FIGURE}") == figure
    assert assistant.find_image(f"/gradio_api/file={quote(str(figure))}") == figure
    assert assistant.find_image("../../../../etc/hosts") is None
    assert assistant.find_image("/etc/hosts") is None and assistant.find_image("nope.png") is None
    shown = served(
        f"Схема:\n![Рис. 1. Инвертор]({MANUAL_FIGURE})\n![Нет](x.png)", assistant.find_image
    )
    assert f"![Рис. 1. Инвертор](/gradio_api/file={quote(str(figure))})" in shown
    assert "*Нет*" in shown and "x.png" not in shown
    remote = "![График](https://example.com/plot.png)"
    assert served(remote, assistant.find_image) == remote


def test_report_uses_manual_figures_and_student_photos(tmp_path):
    def responder(messages):
        if "Составь черновик отчета" in messages[0].text:
            return (
                "## Краткие теоретические сведения\nИнвертор показан на рис. 1.\n\n"
                f"![Рис. 1. Электрическая схема КМОП-пары]({MANUAL_FIGURE})\n\n"
                "## Ход работы\nСобрать схему по рис. 8.\n\n## Выводы\nГотово."
            )
        return course_responder(messages)

    settings = Settings(
        provider="openai", openai_api_key="k", embeddings_provider="none", storage_dir=tmp_path
    )
    llm = ScriptedChatModel(responder=responder)
    assistant = LabAssistant(settings, llm_factory=lambda **kwargs: llm)
    photos = [
        picture(tmp_path / "IMG_0001.jpg"),
        picture(tmp_path / "схема_стенда.png", color="green"),
    ]
    figures = assistant.report_figures([str(p) for p in photos], "Фото установки", Options())
    assert [caption for _, caption in figures] == ["Фото установки", "Схема лабораторной установки"]
    assert all((settings.output_dir / ref).is_file() for ref, _ in figures)
    lab = next(t for t in assistant.lab_titles() if t.startswith("Лабораторная работа 1.4"))
    markdown, path = assistant.make_report(assistant.new_session(), Options(lab=lab), "", figures)
    assert "Инвертор показан на рис. 1." in markdown and "Собрать схему по рис. 2." in markdown
    assert "Рисунок 2. Схема для измерения передаточной (а) и переходной (б)" in markdown
    assert "Рисунок 4. Схема лабораторной" in markdown
    assert len(Document(path).inline_shapes) == 4
