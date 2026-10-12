from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import snowballstemmer
from langchain_core.callbacks import CallbackManagerForRetrieverRun
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.retrievers import BaseRetriever
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from rank_bm25 import BM25Okapi

DOC_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".pptx",
    ".odt",
    ".odp",
    ".rtf",
    ".doc",
    ".html",
    ".htm",
}
CONVERTED = {".rtf", ".doc"}
DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OPEN_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"

_stemmers = {"ru": snowballstemmer.stemmer("russian"), "en": snowballstemmer.stemmer("english")}
_word_re = re.compile(r"\w+", re.UNICODE)
_cyrillic_re = re.compile("[а-яё]")
_lab_number_re = re.compile(
    r"(?:\bлр|работ\w*)\s*[-№]?\s*(?:номер\s*)?(\d+(?:\.\d+)?)", re.IGNORECASE
)
_section_re = re.compile(r"(?m)^(#{1,4})\s+(.+?)\s*$")
OVERVIEW_PARTS = (
    (r"цел[иь]", 2000),
    (r"задани|порядок выполнения|ход работы", 8000),
    (r"отч[её]т", 2000),
    (r"стенд|макетн|оборудован", 1500),
    (r"теорет", 1500),
)


def lab_number(title: str) -> str:
    match = _lab_number_re.search(title)
    return match.group(1) if match else ""


def same_lab(asked: str, number: str) -> bool:
    if not number:
        return False
    return asked == number or ("." not in asked and number.rsplit(".", 1)[-1] == asked)


def sections(text: str) -> list[tuple[int, int, str, str]]:
    heads = [(m.start(), len(m.group(1)), m.group(2)) for m in _section_re.finditer(text)]
    found = []
    for index, (start, level, title) in enumerate(heads):
        end = next((s for s, deeper, _ in heads[index + 1 :] if deeper <= level), len(text))
        found.append((start, level, title, text[start:end].strip()))
    return found


def _clip(text: str, size: int) -> str:
    if len(text) <= size:
        return text
    cut = text.rfind("\n\n", 0, size)
    return text[: cut if cut > size // 2 else size].rstrip() + "\n..."


def overview(text: str, budget: int = 12000) -> str:
    found = [s for s in sections(text) if s[1] >= 2]
    if not found:
        return _clip(text, budget)
    top = min(level for _, level, _, _ in found)
    main = [
        s
        for s in found
        if s[1] == top and not re.search(r"контрольн|литератур|содержани", s[2], re.I)
    ]
    picked, left = {}, budget
    for pattern, cap in OVERVIEW_PARTS:
        for start, _, title, body in main:
            if left <= 0 or start in picked or not re.search(pattern, title, re.I):
                continue
            piece = _clip(body, min(cap, left))
            picked[start], left = piece, left - len(piece)
    return "\n\n".join(picked[start] for start in sorted(picked))


def tokenize(text: str) -> list[str]:
    words = _word_re.findall(text.lower())
    return [_stemmers["ru" if _cyrillic_re.search(w) else "en"].stemWord(w) for w in words]


def decode_text(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace")


def is_text_file(path: Path) -> bool:
    with open(path, "rb") as file:
        head = file.read(4096)
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return True
    if not head or b"\x00" in head:
        return False
    sample = decode_text(head)
    printable = sum(char.isprintable() or char in "\t\n\r\f" for char in sample)
    return printable >= 0.95 * len(sample)


class _HTMLText(HTMLParser):
    BLOCKS = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs) -> None:
        if tag in ("script", "style"):
            self.hidden += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag) -> None:
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data) -> None:
        if not self.hidden:
            self.parts.append(data)


def html_text(source: str) -> str:
    parser = _HTMLText()
    parser.feed(source)
    return re.sub(r"\n\s*\n+", "\n\n", "".join(parser.parts)).strip()


def office_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        if path.suffix.lower() != ".pptx":
            root = ElementTree.fromstring(archive.read("content.xml"))
            blocks = (e for e in root.iter() if e.tag in (OPEN_TEXT + "h", OPEN_TEXT + "p"))
            return "\n".join("".join(block.itertext()) for block in blocks)
        slides = sorted(
            (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"\d+", n).group()),
        )
        parts = []
        for number, name in enumerate(slides, start=1):
            root = ElementTree.fromstring(archive.read(name))
            lines = [
                "".join(t.text or "" for t in p.iter(DRAWING + "t"))
                for p in root.iter(DRAWING + "p")
            ]
            parts.append(f"Слайд {number}\n" + "\n".join(line for line in lines if line.strip()))
        return "\n\n".join(parts)


def converted_text(path: Path) -> str:
    if shutil.which("textutil"):
        command = ["textutil", "-convert", "txt", "-stdout", str(path)]
        result = subprocess.run(command, capture_output=True, timeout=120)
        if result.returncode == 0:
            return decode_text(result.stdout)
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        with tempfile.TemporaryDirectory() as folder:
            profile = (Path(folder) / "profile").as_uri()
            command = [office, f"-env:UserInstallation={profile}", "--headless"]
            command += ["--convert-to", "txt:Text (encoded):UTF8", "--outdir", folder, str(path)]
            subprocess.run(command, capture_output=True, timeout=180)
            for text in Path(folder).glob("*.txt"):
                return decode_text(text.read_bytes())
    raise ValueError(
        f"файлы {path.suffix} читаются в macOS или при установленном LibreOffice; сохраните файл "
        "как DOCX или PDF и загрузите снова"
    )


def load_file(path: str | Path) -> list[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        pages = [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]
        return [
            Document(page_content=text, metadata={"source": path.name, "page": number})
            for number, text in enumerate(pages, start=1)
            if text
        ]
    if ext == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(parts)
    elif ext in (".pptx", ".odt", ".odp"):
        text = office_text(path)
    elif ext in CONVERTED:
        text = converted_text(path)
    elif ext in (".html", ".htm"):
        text = html_text(decode_text(path.read_bytes()))
    elif ext in (".txt", ".md") or is_text_file(path):
        text = decode_text(path.read_bytes())
    else:
        raise ValueError(f"формат {ext or 'без расширения'} не поддерживается")
    return [Document(page_content=text, metadata={"source": path.name})] if text.strip() else []


def format_docs(docs: list[Document]) -> str:
    return "\n\n".join(f"[{d.metadata.get('source', '?')}]\n{d.page_content}" for d in docs)


class HybridRetriever(BaseRetriever):
    kb: Any
    k: int = 4
    only: set[str] | None = None
    prefer: set[str] | None = None
    extra: list[Document] = []

    def _get_relevant_documents(
        self, query: str, *, run_manager: CallbackManagerForRetrieverRun
    ) -> list[Document]:
        return self.extra + self.kb.search(query, self.k, only=self.only, prefer=self.prefer)


class KnowledgeBase:
    def __init__(
        self, embeddings: Embeddings | None = None, chunk_size: int = 900, chunk_overlap: int = 150
    ) -> None:
        self.embeddings = embeddings
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n## ", "\n### ", "\n\n", "\n", ". ", " ", ""],
        )
        self.chunks: list[Document] = []
        self.titles: dict[str, str] = {}
        self.texts: dict[str, str] = {}
        self.vector_store: InMemoryVectorStore | None = None
        self.vector_error = ""
        self._bm25: BM25Okapi | None = None

    def add_directory(self, directory: Path) -> str:
        files = sorted(p for p in Path(directory).glob("*") if p.suffix.lower() in DOC_EXTENSIONS)
        return self.add_files(files) if files else "Каталог методичек пуст"

    def add_files(self, paths: list[str | Path]) -> str:
        report = []
        for path in map(Path, paths):
            try:
                docs = load_file(path)
            except Exception as exc:
                report.append(f"{path.name}: ошибка чтения ({exc})")
                continue
            report.append(self._add(path.name, docs))
        self._rebuild_index()
        return "; ".join(report)

    def add_documents(self, documents: dict[str, list[Document]]) -> str:
        report = [self._add(name, docs) for name, docs in documents.items()]
        self._rebuild_index()
        return "; ".join(report)

    def count(self, name: str) -> int:
        return sum(1 for c in self.chunks if c.metadata["source"] == name)

    def remove(self, names: set[str]) -> None:
        self.chunks = [c for c in self.chunks if c.metadata["source"] not in names]
        for name in names:
            self.titles.pop(name, None)
            self.texts.pop(name, None)
        self._rebuild_index()

    def _add(self, name: str, docs: list[Document]) -> str:
        if not docs:
            return f"{name}: текст не найден"
        self.chunks = [c for c in self.chunks if c.metadata["source"] != name]
        title = self._title(docs[0].page_content, Path(name))
        self.titles[name] = title
        self.texts[name] = "\n\n".join(doc.page_content for doc in docs)
        lab = lab_number(title)
        chunks = self.splitter.split_documents(docs)
        for number, chunk in enumerate(chunks, start=1):
            chunk.metadata.update(source=name, title=title, chunk=number, lab=lab)
            if title not in chunk.page_content[: len(title) + 4]:
                chunk.page_content = f"{title}\n{chunk.page_content}"
        self.chunks += chunks
        return f"{name}: {len(chunks)} фрагм."

    def _rebuild_index(self) -> None:
        self._bm25 = (
            BM25Okapi([tokenize(c.page_content) for c in self.chunks]) if self.chunks else None
        )
        self.vector_store = None
        if self.embeddings is None or not self.chunks:
            return
        try:
            store = InMemoryVectorStore(self.embeddings)
            store.add_documents(self.chunks, ids=[str(i) for i in range(len(self.chunks))])
            self.vector_store, self.vector_error = store, ""
        except Exception as exc:
            self.vector_error = str(exc)

    @staticmethod
    def _title(text: str, path: Path) -> str:
        for line in text.splitlines():
            if line.strip():
                return line.strip("# ").strip()[:200]
        return path.stem

    def search(
        self,
        query: str,
        k: int = 4,
        only: set[str] | None = None,
        prefer: set[str] | None = None,
    ) -> list[Document]:
        if not self.chunks:
            return []
        allowed = [i for i, c in enumerate(self.chunks) if only and c.metadata["source"] in only]
        candidates = allowed or range(len(self.chunks))
        depth = max(3 * k, 10)
        scores = self._bm25.get_scores(tokenize(query))
        bm25_rank = [i for i in sorted(candidates, key=lambda i: -scores[i]) if scores[i] > 0]
        rankings = [bm25_rank[:depth]]
        if allowed:
            rankings.append(allowed[:depth])
        numbers = set(_lab_number_re.findall(query))
        foreign: set[int] = set()
        if numbers:
            wanted = {
                i
                for i in candidates
                if any(same_lab(n, self.chunks[i].metadata.get("lab", "")) for n in numbers)
            }
            if wanted:
                foreign = {i for i in candidates if self.chunks[i].metadata.get("lab")} - wanted
                bm25_rank = [i for i in bm25_rank if i not in foreign]
                rankings[0] = bm25_rank[:depth]
            rankings.append([i for i in bm25_rank if i in wanted] + sorted(wanted - set(bm25_rank)))
        elif prefer:
            chosen = [i for i in candidates if self.chunks[i].metadata["source"] in prefer]
            first = [i for i in bm25_rank if i in set(chosen)]
            rankings.append((first + [i for i in chosen if i not in set(first)])[:depth])
        if self.vector_store is not None:
            keep = (lambda doc: doc.metadata["source"] in only) if allowed else None
            try:
                found = self.vector_store.similarity_search(query, k=depth, filter=keep)
                rankings.append([int(doc.id) for doc in found if int(doc.id) not in foreign])
            except Exception as exc:
                self.vector_error = str(exc)
        fused: dict[int, float] = {}
        for ranking in rankings:
            for rank, i in enumerate(ranking):
                fused[i] = fused.get(i, 0.0) + 1.0 / (60 + rank + 1)
        best = sorted(fused, key=fused.get, reverse=True)[:k]
        return [self.chunks[i] for i in best]

    def as_retriever(
        self,
        k: int = 4,
        only: set[str] | None = None,
        prefer: set[str] | None = None,
        extra: list[Document] | None = None,
    ) -> HybridRetriever:
        return HybridRetriever(kb=self, k=k, only=only, prefer=prefer, extra=extra or [])

    def describe(self) -> str:
        if not self.chunks:
            return "База знаний пуста: загрузите методические указания."
        mode = (
            "векторный + BM25" if self.vector_store is not None else "BM25 (эмбеддинги недоступны)"
        )
        files = "\n".join(f"- {title}" for title in self.titles.values())
        stats = f"Документов: {len(self.titles)}, фрагментов: {len(self.chunks)}, поиск: {mode}"
        return f"{stats}\n\n{files}"
