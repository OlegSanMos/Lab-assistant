from __future__ import annotations

import re
import uuid
from collections.abc import Callable
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from .formulas import TAGS, latex_to_text

FONT = "Times New Roman"
LEFT, CENTER = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
DISPLAY = re.compile(r"\$\$(.+?)\$\$|\\\[(.+?)\\\]", re.S)
TAG = re.compile(r"\\tag\*?\{([^{}]*)\}")
CAPTION = re.compile(
    r"^(Таблица|Рисунок)\s+((?:[А-ЯA-Z]\.)?\d+(?:\.\d+)*)\s*([.:\u2013\u2014-]?)\s*(.*)$"
)
FORMULA_NUMBER = re.compile(r"^\((\d+(?:\.\d+)*)\)$")
HEADING = re.compile(r"^(#{2,4})\s+(.*)")
BULLET = re.compile(r"^\s*[-*]\s+(.*)")
NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)")
RULE = re.compile(r"^\s*([-*_])\1{2,}\s*$")
IMAGE = re.compile(r"!\[([^\]\n]*)\]\(\s*(?:<([^>\n]+)>|([^)\s]+))(?:\s+\"[^\"]*\")?\s*\)")
IMAGE_LINE = re.compile(rf"^\s*{IMAGE.pattern}\s*$")
FIGURE_NAME = re.compile(r"^\s*(?:рис\.|рисунок)\s*(\d+)\s*[.:\u2013\u2014-]?\s*", re.IGNORECASE)
CAPTION_LINE = re.compile(r"^\s*\*?\s*(?:рис\.|рисунок)\s*\d+\s*[.:\u2013\u2014-]", re.IGNORECASE)
FIGURE_REF = re.compile(r"(?<![а-яё])(рис\.|рисун[а-яё]*)(\s*)(\d+)", re.IGNORECASE)
DOCX_IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
PLOT_NAME = "Экспериментальная зависимость и аппроксимация МНК"


def _add_text(paragraph, text: str, bold: bool = False) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|<su[bp]>.*?</su[bp]>)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _add_text(paragraph, part[2:-2], bold=True)
        elif part:
            run = paragraph.add_run(TAGS.sub("", part))
            run.bold = bold or None
            if part.startswith("<sub>"):
                run.font.subscript = True
            elif part.startswith("<sup>"):
                run.font.superscript = True


def _plain_font(rpr) -> None:
    fonts = rpr.get_or_add_rFonts()
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{name}"), None)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)


def _setup(document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    _plain_font(document.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr")[0])
    normal = document.styles["Normal"]
    _plain_font(normal.element.get_or_add_rPr())
    normal.font.size = Pt(14)
    fmt = normal.paragraph_format
    fmt.line_spacing, fmt.first_line_indent = 1.5, Cm(1.25)
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        _plain_font(style.element.get_or_add_rPr())
        style.font.size, style.font.bold, style.font.italic = Pt(14), True, False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.alignment = LEFT


def _number(value) -> str:
    if pd.isna(value):
        return ""
    return f"{value:.6g}".replace(".", ",") if isinstance(value, float) else str(value)


def _has_data(table: pd.DataFrame | None, plot: str | None) -> tuple[bool, bool]:
    return table is not None and len(table) > 0, bool(plot) and Path(plot).exists()


def report_inserts(
    table: pd.DataFrame | None, plot: str | None, figures: list[tuple[str, str]] = ()
) -> str:
    has_table, has_plot = _has_data(table, plot)
    parts = []
    items = ["таблицу 1 с результатами измерений"] if has_table else []
    if has_plot:
        items.append("график зависимости")
    if items:
        parts.append(
            "программа сама вставит в начало раздела «Результаты измерений и их обработка» "
            + " и ".join(items)
            + " и сошлется на них, не повторяй эти данные"
        )
    if figures:
        listed = "; ".join(f"![{caption}]({ref})" for ref, caption in figures)
        parts.append(
            "рисунки студента вставь отдельными строками в подходящие разделы, не меняя путь: "
            + listed
        )
    parts.append("свои таблицы нумеруй " + ("со 2" if has_table else "с 1"))
    return "; ".join(parts) + "."


def _existing(ref: str) -> Path | None:
    path = Path(ref)
    return path if path.is_absolute() and path.is_file() else None


def _figure(name: str, path: Path | None, old: int | None = None) -> dict:
    return {"kind": "figure", "name": name, "path": path, "old": old, "number": 0}


def _measurement_blocks(table, plot, extra: list[dict]) -> list:
    has_table, has_plot = _has_data(table, plot)
    blocks = []
    if has_table:
        labels = table.attrs.get("labels", {})
        header = [str(labels.get(column, column)) for column in table.columns]
        rows = [[_number(value) for value in row] for row in table.head(50).itertuples(index=False)]
        blocks += [
            "Результаты измерений приведены в таблице 1.",
            "",
            "Таблица 1 – Результаты измерений",
        ]
        blocks += ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
        blocks += ["| " + " | ".join(row) + " |" for row in rows] + [""]
    if has_plot:
        plot_figure = _figure(PLOT_NAME, Path(plot))
        blocks += [
            {
                "kind": "ref",
                "figures": [plot_figure],
                "text": "Экспериментальная зависимость и ее аппроксимация по методу наименьших "
                "квадратов показаны на {}.",
            },
            plot_figure,
        ]
    if extra:
        text = "Материалы, приложенные к работе, приведены на {}."
        blocks += [{"kind": "ref", "figures": extra, "text": text}, *extra]
    return blocks


def _figure_words(numbers: list[int]) -> str:
    if len(numbers) == 1:
        return f"рисунке {numbers[0]}"
    return "рисунках " + ", ".join(map(str, numbers[:-1])) + f" и {numbers[-1]}"


def figure_catalog(text: str) -> dict[int, tuple[str, str]]:
    catalog = {}
    for match in IMAGE.finditer(text):
        number = FIGURE_NAME.match(match.group(1))
        if number:
            name = FIGURE_NAME.sub("", match.group(1)).strip(" *.")
            catalog.setdefault(int(number.group(1)), (name, match.group(2) or match.group(3)))
    return catalog


def _with_manual_figures(blocks: list, manual: dict, find: Callable) -> list:
    placed = {b["old"] for b in blocks if isinstance(b, dict) and b["old"] is not None}
    result, waiting = [], []
    for index, block in enumerate(blocks):
        if isinstance(block, str) and CAPTION_LINE.match(block):
            number = int(FIGURE_NAME.match(block.replace("*", "")).group(1))
            if number in manual and number not in placed:
                placed.add(number)
                result.append(_figure(manual[number][0], find(manual[number][1]), number))
                continue
        result.append(block)
        if isinstance(block, str):
            for match in FIGURE_REF.finditer(block):
                number = int(match.group(3))
                if number in manual and number not in placed:
                    path = find(manual[number][1])
                    if path is not None:
                        placed.add(number)
                        waiting.append(_figure(manual[number][0], path, number))
        following = blocks[index + 1] if index + 1 < len(blocks) else ""
        ended = not isinstance(following, str) or not following.strip()
        if waiting and (ended or HEADING.match(following)):
            result += ["", *waiting, ""]
            waiting = []
    return result


def compose_report(
    markdown: str,
    table: pd.DataFrame | None = None,
    plot: str | None = None,
    figures: list[tuple[str, str]] = (),
    find_image: Callable[[str], Path | None] | None = None,
    manual: dict[int, tuple[str, str]] | None = None,
) -> str:
    find = find_image or _existing
    blocks: list = []
    caption_may_follow = False
    for line in IMAGE.sub(lambda m: f"\n{m.group(0)}\n", markdown).splitlines():
        image = IMAGE_LINE.match(line)
        if image:
            old = FIGURE_NAME.match(image.group(1))
            name = FIGURE_NAME.sub("", image.group(1)).strip(" *.")
            ref = image.group(2) or image.group(3)
            blocks.append(_figure(name, find(ref), int(old.group(1)) if old else None))
            caption_may_follow = True
            continue
        if caption_may_follow and CAPTION_LINE.match(line):
            figure = next(b for b in reversed(blocks) if isinstance(b, dict))
            figure["name"] = figure["name"] or FIGURE_NAME.sub("", line.replace("*", "")).strip(
                " ."
            )
            caption_may_follow = False
            continue
        caption_may_follow = caption_may_follow and not line.strip()
        blocks.append(line)
    if manual:
        blocks = _with_manual_figures(blocks, manual, find)
    placed = {b["path"] for b in blocks if isinstance(b, dict) and b["path"]}
    extra = [
        _figure(caption, path)
        for ref, caption in figures
        if (path := find(ref)) is not None and path not in placed
    ]
    inserts = _measurement_blocks(table, plot, extra)
    if inserts:
        results = next(
            (
                i
                for i, b in enumerate(blocks)
                if isinstance(b, str) and HEADING.match(b) and "результат" in b.lower()
            ),
            None,
        )
        if results is None:
            results = next(
                (
                    i
                    for i, b in enumerate(blocks)
                    if isinstance(b, str) and HEADING.match(b) and "вывод" in b.lower()
                ),
                len(blocks),
            )
            blocks[results:results] = ["## Результаты измерений и их обработка", ""]
        blocks[results + 1 : results + 1] = ["", *inserts]
    mapping, count = {}, 0
    for block in blocks:
        if isinstance(block, dict) and block["kind"] == "figure" and block["path"]:
            count += 1
            block["number"] = count
            if block["old"] is not None:
                mapping.setdefault(block["old"], count)
    output = []
    for block in blocks:
        if isinstance(block, str):
            output.append(
                FIGURE_REF.sub(
                    lambda m: (
                        m.group(1) + m.group(2) + str(mapping.get(int(m.group(3)), m.group(3)))
                    ),
                    block,
                )
            )
        elif block["kind"] == "ref":
            output.append(
                block["text"].format(_figure_words([f["number"] for f in block["figures"]]))
            )
        elif block["path"]:
            output += ["", f"![Рисунок {block['number']}. {block['name']}](<{block['path']}>)", ""]
        else:
            output.append(f"[вставить рисунок: {block['name'] or 'рисунок из методички'}]")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(output)).strip() + "\n"


class _Writer:
    def __init__(self, document, folder: Path) -> None:
        self.doc, self.folder = document, folder
        self.tables = self.formulas = 0
        self.after_blank = self.captioned = False

    def add(self, text: str = "", indent: bool = True, align=None, bold: bool = False):
        paragraph = self.doc.add_paragraph()
        if not indent:
            paragraph.paragraph_format.first_line_indent = Cm(0)
        if align is not None:
            paragraph.paragraph_format.alignment = align
        _add_text(paragraph, text, bold)
        self.after_blank = not text
        return paragraph

    def blank(self) -> None:
        if not self.after_blank:
            self.add()

    def heading(self, text: str, level: int) -> None:
        text = TAGS.sub("", text).replace("*", "").strip().rstrip(".").strip()
        if not text:
            return
        if level == 1:
            self.blank()
        self.doc.add_heading(text[0].upper() + text[1:], level=level)
        self.after_blank = False
        if level == 1:
            self.blank()

    def caption(self, kind: str, number: str, name: str) -> None:
        if kind == "Таблица":
            text = f"Таблица {number} – {name}" if name else f"Таблица {number}"
            self.add(text, indent=False, align=LEFT).paragraph_format.keep_with_next = True
            if number.isdigit():
                self.tables = max(self.tables, int(number))
            self.captioned = True
        else:
            self.add(f"Рисунок {number}. {name}" if name else f"Рисунок {number}", False, CENTER)

    def table(self, rows: list[list[str]]) -> None:
        if not self.captioned:
            self.caption("Таблица", str(self.tables + 1), "")
        width = max(len(row) for row in rows)
        grid = self.doc.add_table(rows=len(rows), cols=width)
        grid.style = "Table Grid"
        for row, values in zip(grid.rows, rows, strict=True):
            for cell, value in zip(row.cells, values + [""] * (width - len(values)), strict=True):
                fmt = cell.paragraphs[0].paragraph_format
                fmt.first_line_indent, fmt.line_spacing, fmt.alignment = Cm(0), 1.0, CENTER
                _add_text(cell.paragraphs[0], value)
        grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        self.after_blank = self.captioned = False

    def markdown_table(self, lines: list[str]) -> None:
        rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
        rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
        if rows:
            self.table(rows)

    def formula(self, source: str, number: str = "") -> None:
        tag = TAG.search(source)
        number = tag.group(1).strip() if tag else number
        if number.isdigit():
            self.formulas = int(number)
        elif not number:
            self.formulas += 1
            number = str(self.formulas)
        text = latex_to_text(f"$${TAG.sub('', source)}$$", markup=True).strip()
        paragraph = self.add(indent=False, align=LEFT)
        stops = paragraph.paragraph_format.tab_stops
        stops.add_tab_stop(Cm(8.25), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
        _add_text(paragraph, f"\t{text}\t({number})")
        self.after_blank = False

    def picture(self, path: str, caption: str) -> None:
        file = Path(path)
        try:
            with Image.open(file) as image:
                width, height = image.size
                if file.suffix.lower() not in DOCX_IMAGES:
                    self.folder.mkdir(parents=True, exist_ok=True)
                    file = self.folder / f"{file.stem}_{uuid.uuid4().hex[:6]}.png"
                    image.convert("RGB").save(file)
        except Exception:
            self.add(f"[не удалось вставить рисунок: {caption}]")
            return
        paragraph = self.add(indent=False, align=CENTER)
        paragraph.add_run().add_picture(str(file), width=Cm(min(15.0, 12.0 * width / height)))
        paragraph.paragraph_format.keep_with_next = True
        match = CAPTION.match(caption)
        if match:
            self.caption("Рисунок", match.group(2), match.group(4).strip())
        elif caption:
            self.add(caption, indent=False, align=CENTER)


def _items(markdown: str) -> list[tuple[str, str]]:
    items, start = [], 0
    for match in DISPLAY.finditer(markdown):
        text = latex_to_text(markdown[start : match.start()], markup=True)
        items += [("line", line) for line in text.splitlines()]
        items.append(("formula", match.group(1) or match.group(2)))
        start = match.end()
    text = latex_to_text(markdown[start:], markup=True)
    return items + [("line", line) for line in text.splitlines()]


def _next_line(items: list[tuple[str, str]], index: int) -> str:
    for kind, line in items[index:]:
        if kind != "line":
            return ""
        if line.strip():
            return line.strip()
    return ""


def _formula_number(items: list[tuple[str, str]], index: int) -> tuple[str, int]:
    for position in range(index, len(items)):
        kind, line = items[position]
        if kind == "line" and not line.strip():
            continue
        match = FORMULA_NUMBER.match(line.strip()) if kind == "line" else None
        return (match.group(1), position + 1) if match else ("", index)
    return "", index


def _caption(line: str, next_line: str) -> re.Match | None:
    match = CAPTION.match(line.replace("*", "").strip())
    if not match or not (match.group(3) or not match.group(4)):
        return None
    if match.group(1) == "Таблица" and not next_line.startswith("|"):
        return None
    return match


def _is_row(item: tuple[str, str]) -> bool:
    return item[0] == "line" and item[1].lstrip().startswith("|")


def export_report_docx(
    markdown: str,
    title: str,
    out_dir: Path,
    table: pd.DataFrame | None = None,
    plot: str | None = None,
    figures: list[tuple[str, str]] = (),
    find_image: Callable[[str], Path | None] | None = None,
    manual: dict[int, tuple[str, str]] | None = None,
) -> str:
    final = compose_report(markdown, table, plot, figures, find_image, manual)
    return render_report_docx(final, title, out_dir)


def render_report_docx(markdown: str, title: str, out_dir: Path) -> str:
    document = Document()
    _setup(document)
    writer = _Writer(document, out_dir)
    writer.add("Отчет по лабораторной работе", indent=False, align=CENTER, bold=True)
    if title:
        writer.add(title, indent=False, align=CENTER, bold=True)
    items, index = _items(markdown), 0
    while index < len(items):
        kind, line = items[index]
        index += 1
        if kind == "formula":
            number, index = _formula_number(items, index)
            writer.formula(line, number)
            continue
        line = line.rstrip()
        if not line.strip() or line.startswith("# ") or RULE.match(line):
            continue
        image = IMAGE_LINE.match(line)
        if image:
            writer.picture(image.group(2) or image.group(3), image.group(1))
            continue
        if _is_row((kind, line)):
            rows = [line]
            while index < len(items) and _is_row(items[index]):
                rows.append(items[index][1])
                index += 1
            writer.markdown_table(rows)
            continue
        heading, caption = HEADING.match(line), _caption(line, _next_line(items, index))
        bullet, number = BULLET.match(line), NUMBERED.match(line)
        if heading:
            writer.heading(TAGS.sub("", heading.group(2)).strip(), len(heading.group(1)) - 1)
        elif caption:
            writer.caption(caption.group(1), caption.group(2), caption.group(4).strip())
        elif bullet:
            _add_text(document.add_paragraph(style="List Bullet"), bullet.group(1))
            writer.after_blank = False
        elif number:
            _add_text(document.add_paragraph(style="List Number"), number.group(1))
            writer.after_blank = False
        elif re.match(r"где\b", line.strip()):
            writer.add(re.sub(r"^где\s*:?\s*", "где ", line.strip()), indent=False)
        else:
            writer.add(line)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"report_{uuid.uuid4().hex[:8]}.docx"
    document.save(str(path))
    return str(path)
